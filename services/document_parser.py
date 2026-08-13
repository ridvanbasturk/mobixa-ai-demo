"""Yüklenen dosyalardan metin çıkarma (TXT, PDF, DOCX, PPTX)."""
from __future__ import annotations

from dataclasses import dataclass
from typing import Optional


@dataclass
class ParseResult:
    success: bool
    text: Optional[str]
    error: Optional[str]


def extract_text_from_txt(file_bytes: bytes) -> ParseResult:
    try:
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text = file_bytes.decode("utf-8", errors="replace")
        return ParseResult(True, text, None)
    except Exception as exc:  # noqa: BLE001
        return ParseResult(False, None, f"TXT dosyası okunamadı: {exc}")


def extract_text_from_pdf(file_bytes: bytes) -> ParseResult:
    try:
        from pypdf import PdfReader
        import io

        reader = PdfReader(io.BytesIO(file_bytes))
        pages = [page.extract_text() or "" for page in reader.pages]
        text = "\n".join(pages).strip()
        if not text:
            return ParseResult(False, None, "PDF içinden metin çıkarılamadı (taranmış görsel olabilir)")
        return ParseResult(True, text, None)
    except Exception as exc:  # noqa: BLE001
        return ParseResult(False, None, f"PDF dosyası okunamadı: {exc}")


def extract_text_from_docx(file_bytes: bytes) -> ParseResult:
    try:
        from docx import Document
        import io

        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        text = "\n".join(paragraphs).strip()
        if not text:
            return ParseResult(False, None, "DOCX içinden metin çıkarılamadı")
        return ParseResult(True, text, None)
    except Exception as exc:  # noqa: BLE001
        return ParseResult(False, None, f"DOCX dosyası okunamadı: {exc}")


def extract_text_from_pptx(file_bytes: bytes) -> ParseResult:
    try:
        from pptx import Presentation
        import io

        prs = Presentation(io.BytesIO(file_bytes))
        chunks = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    chunks.append(shape.text)
        text = "\n".join(chunks).strip()
        if not text:
            return ParseResult(False, None, "PPTX içinden metin çıkarılamadı")
        return ParseResult(True, text, None)
    except Exception as exc:  # noqa: BLE001
        return ParseResult(False, None, f"PPTX dosyası okunamadı: {exc}")


def extract_text(file_bytes: bytes, filename: str) -> ParseResult:
    """Dosya uzantısına göre uygun çıkarıcıyı seçer."""
    lower = filename.lower()
    if lower.endswith(".txt"):
        return extract_text_from_txt(file_bytes)
    if lower.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    if lower.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    if lower.endswith(".pptx"):
        return extract_text_from_pptx(file_bytes)
    return ParseResult(False, None, f"Desteklenmeyen dosya türü: {filename}")
